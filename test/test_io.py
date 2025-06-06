import os
from pathlib import Path

import pandas as pd
import pytest
from geopandas.testing import assert_geodataframe_equal
from numpy.testing import assert_array_almost_equal as equal

import pypsa

try:
    import tables  # noqa: F401

    tables_installed = True
except ImportError:
    tables_installed = False

try:
    import openpyxl  # noqa: F401
    import python_calamine  # noqa: F401

    excel_installed = True
except ImportError:
    excel_installed = False


# TODO classes could be further parametrized
class TestCSVDir:
    @pytest.mark.parametrize(
        "meta",
        [
            {"test": "test"},
            {"test": "test", "test2": "test2"},
            {"test": {"test": "test", "test2": "test2"}},
        ],
    )
    def test_csv_io(self, scipy_network, tmpdir, meta):
        fn = os.path.join(tmpdir, "csv_export")
        scipy_network.meta = meta
        scipy_network.export_to_csv_folder(fn)
        pypsa.Network(fn)
        reloaded = pypsa.Network(fn)
        assert reloaded.meta == scipy_network.meta

    @pytest.mark.parametrize(
        "meta",
        [
            {"test": "test"},
            {"test": "test", "test2": "test2"},
            {"test": {"test": "test", "test2": "test2"}},
        ],
    )
    def test_csv_io_quotes(self, scipy_network, tmpdir, meta, quotechar="'"):
        fn = os.path.join(tmpdir, "csv_export")
        scipy_network.meta = meta
        scipy_network.export_to_csv_folder(fn, quotechar=quotechar)
        imported = pypsa.Network()
        imported.import_from_csv_folder(fn, quotechar=quotechar)
        assert imported.meta == scipy_network.meta

    def test_csv_io_Path(self, scipy_network, tmpdir):
        fn = Path(os.path.join(tmpdir, "csv_export"))
        scipy_network.export_to_csv_folder(fn)
        pypsa.Network(fn)

    def test_csv_io_multiindexed(self, ac_dc_network_mi, tmpdir):
        fn = os.path.join(tmpdir, "csv_export")
        ac_dc_network_mi.export_to_csv_folder(fn)
        m = pypsa.Network(fn)
        pd.testing.assert_frame_equal(
            m.generators_t.p,
            ac_dc_network_mi.generators_t.p,
        )

    def test_csv_io_shapes(self, ac_dc_network_shapes, tmpdir):
        fn = os.path.join(tmpdir, "csv_export")
        ac_dc_network_shapes.export_to_csv_folder(fn)
        m = pypsa.Network(fn)
        assert_geodataframe_equal(
            m.shapes,
            ac_dc_network_shapes.shapes,
            check_less_precise=True,
        )

    def test_csv_io_shapes_with_missing(self, ac_dc_network_shapes, tmpdir):
        fn = os.path.join(tmpdir, "csv_export")
        n = ac_dc_network_shapes.copy()
        n.shapes.loc["Manchester", "geometry"] = None
        n.export_to_csv_folder(fn)
        m = pypsa.Network(fn)
        assert_geodataframe_equal(
            m.shapes,
            n.shapes,
            check_less_precise=True,
        )


class TestNetcdf:
    @pytest.mark.parametrize(
        "meta",
        [
            {"test": "test"},
            {"test": "test", "test2": "test2"},
            {"test": {"test": "test", "test2": "test2"}},
        ],
    )
    def test_netcdf_io(self, scipy_network, tmpdir, meta):
        fn = os.path.join(tmpdir, "netcdf_export.nc")
        scipy_network.meta = meta
        scipy_network.export_to_netcdf(fn)
        reloaded = pypsa.Network(fn)
        assert reloaded.meta == scipy_network.meta

    def test_netcdf_io_Path(self, scipy_network, tmpdir):
        fn = Path(os.path.join(tmpdir, "netcdf_export.nc"))
        scipy_network.export_to_netcdf(fn)
        pypsa.Network(fn)

    def test_netcdf_io_datetime(self, tmpdir):
        fn = os.path.join(tmpdir, "temp.nc")
        exported_sns = pd.date_range(start="2013-03-01", end="2013-03-02", freq="h")
        n = pypsa.Network()
        n.set_snapshots(exported_sns)
        n.export_to_netcdf(fn)
        imported_sns = pypsa.Network(fn).snapshots

        assert (imported_sns == exported_sns).all()

    def test_netcdf_io_multiindexed(self, ac_dc_network_mi, tmpdir):
        fn = os.path.join(tmpdir, "netcdf_export.nc")
        ac_dc_network_mi.export_to_netcdf(fn)
        m = pypsa.Network(fn)
        pd.testing.assert_frame_equal(
            m.generators_t.p,
            ac_dc_network_mi.generators_t.p,
        )
        pd.testing.assert_frame_equal(
            m.snapshot_weightings,
            ac_dc_network_mi.snapshot_weightings[
                m.snapshot_weightings.columns
            ],  # reset order
        )

    def test_netcdf_io_shapes(self, ac_dc_network_shapes, tmpdir):
        fn = os.path.join(tmpdir, "netcdf_export.nc")
        ac_dc_network_shapes.export_to_netcdf(fn)
        m = pypsa.Network(fn)
        assert_geodataframe_equal(
            m.shapes,
            ac_dc_network_shapes.shapes,
            check_less_precise=True,
        )

    def test_netcdf_io_shapes_with_missing(self, ac_dc_network_shapes, tmpdir):
        fn = os.path.join(tmpdir, "netcdf_export.nc")
        n = ac_dc_network_shapes.copy()
        n.shapes.loc["Manchester", "geometry"] = None
        n.export_to_netcdf(fn)
        m = pypsa.Network(fn)
        assert_geodataframe_equal(
            m.shapes,
            n.shapes,
            check_less_precise=True,
        )

    def test_netcdf_from_url(self):
        url = "https://github.com/PyPSA/PyPSA/raw/master/examples/scigrid-de/scigrid-with-load-gen-trafos.nc"
        pypsa.Network(url)

    def test_netcdf_io_no_compression(self, scipy_network, tmpdir):
        fn = os.path.join(tmpdir, "netcdf_export.nc")
        scipy_network.export_to_netcdf(fn, float32=False, compression=None)
        scipy_network_compressed = pypsa.Network(fn)
        assert (
            (scipy_network.loads_t.p_set == scipy_network_compressed.loads_t.p_set)
            .all()
            .all()
        )

    def test_netcdf_io_custom_compression(self, scipy_network, tmpdir):
        fn = os.path.join(tmpdir, "netcdf_export.nc")
        digits = 5
        compression = dict(zlib=True, complevel=9, least_significant_digit=digits)
        scipy_network.export_to_netcdf(fn, compression=compression)
        scipy_network_compressed = pypsa.Network(fn)
        assert (
            (
                (
                    scipy_network.loads_t.p_set - scipy_network_compressed.loads_t.p_set
                ).abs()
                < 1 / 10**digits
            )
            .all()
            .all()
        )

    def test_netcdf_io_typecast(self, scipy_network, tmpdir):
        fn = os.path.join(tmpdir, "netcdf_export.nc")
        scipy_network.export_to_netcdf(fn, float32=True, compression=None)
        pypsa.Network(fn)

    def test_netcdf_io_typecast_and_compression(self, scipy_network, tmpdir):
        fn = os.path.join(tmpdir, "netcdf_export.nc")
        scipy_network.export_to_netcdf(fn, float32=True)
        pypsa.Network(fn)


@pytest.mark.skipif(not tables_installed, reason="PyTables not installed")
class TestHDF5:
    @pytest.mark.parametrize(
        "meta",
        [
            {"test": "test"},
            {"test": "test", "test2": "test2"},
            {"test": {"test": "test", "test2": "test2"}},
        ],
    )
    def test_hdf5_io(self, scipy_network, tmpdir, meta):
        fn = os.path.join(tmpdir, "hdf5_export.h5")
        scipy_network.meta = meta
        scipy_network.export_to_hdf5(fn)
        pypsa.Network(fn)
        reloaded = pypsa.Network(fn)
        assert reloaded.meta == scipy_network.meta

    def test_hdf5_io_Path(self, scipy_network, tmpdir):
        fn = Path(os.path.join(tmpdir, "hdf5_export.h5"))
        scipy_network.export_to_hdf5(fn)
        pypsa.Network(fn)

    def test_hdf5_io_multiindexed(self, ac_dc_network_mi, tmpdir):
        fn = os.path.join(tmpdir, "hdf5_export.h5")
        ac_dc_network_mi.export_to_hdf5(fn)
        m = pypsa.Network(fn)
        pd.testing.assert_frame_equal(
            m.generators_t.p,
            ac_dc_network_mi.generators_t.p,
        )

    def test_hdf5_io_shapes(self, ac_dc_network_shapes, tmpdir):
        fn = os.path.join(tmpdir, "hdf5_export.h5")
        ac_dc_network_shapes.export_to_hdf5(fn)
        m = pypsa.Network(fn)
        assert_geodataframe_equal(
            m.shapes,
            ac_dc_network_shapes.shapes,
            check_less_precise=True,
        )

    def test_hdf5_io_shapes_with_missing(self, ac_dc_network_shapes, tmpdir):
        fn = os.path.join(tmpdir, "hdf5_export.h5")
        n = ac_dc_network_shapes.copy()
        n.shapes.loc["Manchester", "geometry"] = None
        n.export_to_hdf5(fn)
        m = pypsa.Network(fn)
        assert_geodataframe_equal(
            m.shapes,
            n.shapes,
            check_less_precise=True,
        )


@pytest.mark.skipif(not excel_installed, reason="openpyxl not installed")
class TestExcelIO:
    @pytest.mark.parametrize(
        "meta",
        [
            {"test": "test"},
            {"test": "test", "test2": "test2"},
            {"test": {"test": "test", "test2": "test2"}},
        ],
    )
    def test_excel_io(self, scipy_network, tmpdir, meta):
        fn = os.path.join(tmpdir, "excel_export.xlsx")
        scipy_network.meta = meta
        scipy_network.export_to_excel(fn)
        reloaded = pypsa.Network(fn)
        assert reloaded.meta == scipy_network.meta

    def test_excel_io_Path(self, scipy_network, tmpdir):
        fn = Path(os.path.join(tmpdir, "excel_export.xlsx"))
        scipy_network.export_to_excel(fn)
        pypsa.Network(fn)

    def test_excel_io_datetime(self, tmpdir):
        fn = os.path.join(tmpdir, "temp.xlsx")
        exported_sns = pd.date_range(start="2013-03-01", end="2013-03-02", freq="h")
        n = pypsa.Network()
        n.set_snapshots(exported_sns)
        n.export_to_excel(fn)
        imported_sns = pypsa.Network(fn).snapshots
        assert (imported_sns == exported_sns).all()

    def test_excel_io_multiindexed(self, ac_dc_network_mi, tmpdir):
        fn = os.path.join(tmpdir, "excel_export.xlsx")
        ac_dc_network_mi.export_to_excel(fn)
        m = pypsa.Network(fn)
        pd.testing.assert_frame_equal(
            m.generators_t.p,
            ac_dc_network_mi.generators_t.p,
        )
        pd.testing.assert_frame_equal(
            m.snapshot_weightings,
            ac_dc_network_mi.snapshot_weightings[m.snapshot_weightings.columns],
            check_dtype=False,  # TODO Remove once validation layer leads to safer types
        )

    def test_excel_io_shapes(self, ac_dc_network_shapes, tmpdir):
        fn = os.path.join(tmpdir, "excel_export.xlsx")
        ac_dc_network_shapes.export_to_excel(fn)
        m = pypsa.Network(fn)
        assert_geodataframe_equal(
            m.shapes,
            ac_dc_network_shapes.shapes,
            check_less_precise=True,
        )

    def test_excel_io_shapes_with_missing(self, ac_dc_network_shapes, tmpdir):
        fn = os.path.join(tmpdir, "excel_export.xlsx")
        n = ac_dc_network_shapes.copy()
        n.shapes.loc["Manchester", "geometry"] = None
        n.export_to_excel(fn)
        m = pypsa.Network(fn)
        assert_geodataframe_equal(
            m.shapes,
            n.shapes,
            check_less_precise=True,
        )

    def test_io_time_dependent_efficiencies_excel(self, tmpdir):
        n = pypsa.Network()
        s = [1, 0.95, 0.99]
        n.snapshots = range(len(s))
        n.add("Bus", "bus")
        n.add("Generator", "gen", bus="bus", efficiency=s)
        n.add("Store", "sto", bus="bus", standing_loss=s)
        n.add(
            "StorageUnit",
            "su",
            bus="bus",
            efficiency_store=s,
            efficiency_dispatch=s,
            standing_loss=s,
        )
        fn = os.path.join(tmpdir, "network-time-eff.xlsx")
        n.export_to_excel(fn)
        m = pypsa.Network(fn)
        assert not m.stores_t.standing_loss.empty
        assert not m.storage_units_t.standing_loss.empty
        assert not m.generators_t.efficiency.empty
        assert not m.storage_units_t.efficiency_store.empty
        assert not m.storage_units_t.efficiency_dispatch.empty
        equal(m.stores_t.standing_loss, n.stores_t.standing_loss)
        equal(m.storage_units_t.standing_loss, n.storage_units_t.standing_loss)
        equal(m.generators_t.efficiency, n.generators_t.efficiency)
        equal(m.storage_units_t.efficiency_store, n.storage_units_t.efficiency_store)
        equal(
            m.storage_units_t.efficiency_dispatch, n.storage_units_t.efficiency_dispatch
        )


@pytest.mark.parametrize("use_pandapower_index", [True, False])
@pytest.mark.parametrize("extra_line_data", [True, False])
def test_import_from_pandapower_network(
    pandapower_custom_network,
    pandapower_cigre_network,
    extra_line_data,
    use_pandapower_index,
):
    nets = [pandapower_custom_network, pandapower_cigre_network]
    for net in nets:
        n = pypsa.Network()
        n.import_from_pandapower_net(
            net,
            use_pandapower_index=use_pandapower_index,
            extra_line_data=extra_line_data,
        )
        assert len(n.buses) == len(net.bus)
        assert len(n.generators) == (len(net.gen) + len(net.sgen) + len(net.ext_grid))
        assert len(n.loads) == len(net.load)
        assert len(n.transformers) == len(net.trafo)
        assert len(n.shunt_impedances) == len(net.shunt)


def test_io_time_dependent_efficiencies(tmpdir):
    n = pypsa.Network()
    s = [1, 0.95, 0.99]
    n.snapshots = range(len(s))
    n.add("Bus", "bus")
    n.add("Generator", "gen", bus="bus", efficiency=s)
    n.add("Store", "sto", bus="bus", standing_loss=s)
    n.add(
        "StorageUnit",
        "su",
        bus="bus",
        efficiency_store=s,
        efficiency_dispatch=s,
        standing_loss=s,
    )

    fn = os.path.join(tmpdir, "network-time-eff.nc")
    n.export_to_netcdf(fn)
    m = pypsa.Network(fn)

    assert not m.stores_t.standing_loss.empty
    assert not m.storage_units_t.standing_loss.empty
    assert not m.generators_t.efficiency.empty
    assert not m.storage_units_t.efficiency_store.empty
    assert not m.storage_units_t.efficiency_dispatch.empty

    equal(m.stores_t.standing_loss, n.stores_t.standing_loss)
    equal(m.storage_units_t.standing_loss, n.storage_units_t.standing_loss)
    equal(m.generators_t.efficiency, n.generators_t.efficiency)
    equal(m.storage_units_t.efficiency_store, n.storage_units_t.efficiency_store)
    equal(m.storage_units_t.efficiency_dispatch, n.storage_units_t.efficiency_dispatch)
